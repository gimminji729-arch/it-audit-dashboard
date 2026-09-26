"""
build_rcm_excel.py
------------------------------------------------------------
Generates docs/Risk_and_Control_Matrix.xlsx from the single
source of truth in logic.RISK_CONTROL_MATRIX, so the workbook
and the in-app Risk & Control Matrix page never drift apart.

Run:  python build_rcm_excel.py
------------------------------------------------------------
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import logic

OUT_PATH = Path(__file__).parent / "docs" / "Risk_and_Control_Matrix.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F3864")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DOMAIN_FILL = {
    "Access Management": PatternFill("solid", fgColor="E2EFDA"),
    "Change Management": PatternFill("solid", fgColor="FFF2CC"),
}

COLUMNS = [
    "Domain",
    "Risk",
    "Control Objective",
    "Control Activity",
    "Control Type",
    "Frequency",
    "Test Procedure",
    "SQL Test",
    "Framework Mapping",
]
COLUMN_WIDTHS = [16, 34, 32, 40, 16, 14, 42, 30, 40]


def build():
    wb = Workbook()

    notes = wb.active
    notes.title = "Read Me"
    _build_notes_sheet(notes)

    ws = wb.create_sheet("Risk and Control Matrix")
    _build_rcm_sheet(ws)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


def _build_notes_sheet(ws: Worksheet):
    ws.column_dimensions["A"].width = 100
    ws["A1"] = f"{logic.COMPANY_NAME} — ITGC Risk and Control Matrix"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Access Management & Change Management controls | Audit as-of date: " + logic.AS_OF_DATE
    ws["A2"].font = SUBTITLE_FONT

    lines = [
        "",
        "About this workbook",
        "This Risk and Control Matrix (RCM) documents the IT general controls (ITGC) tested "
        "in the accompanying User Access & Change Management Audit Dashboard project.",
        "",
        "IMPORTANT: Hankuk Motors America, Inc. is a fictional company created for this "
        "portfolio project. All employees, accounts, roles, and change records are "
        "synthetic and were generated for demonstration purposes only. No real company, "
        "employee, or system is represented.",
        "",
        "How to read this matrix",
        "Each row pairs one risk with the control designed to address it, the procedure an "
        "auditor would perform to test that control, the SQL script in this project that "
        "performs that test, and the external framework(s) the control maps to.",
        "",
        "Frameworks referenced",
        "- PCAOB Auditing Standard No. 2201 (AS 2201) and the COSO Internal Control - "
        "Integrated Framework (Control Activities component), specifically the ITGC domains "
        "'Access to Programs and Data' and 'Program Change Management'.",
        "- 내부회계관리제도 모범규준 (K-ICFR), the Korean Internal Accounting Control System "
        "standard, which parallels US SOX/COSO and applies to the Korea-listed parent "
        "company of the fictional US subsidiary used in this scenario.",
        "",
        "Related files in this project",
        "- /sql — the SQL control test scripts referenced in the 'SQL Test' column",
        "- /docs/AUDIT_PROGRAM.md — audit objective, scope, population, and risk rating methodology",
        "- /docs/SAMPLE_AUDIT_REPORT.md — example audit report built from this cycle's results",
        "- app.py — interactive dashboard that runs these tests and tracks findings to remediation",
    ]
    row = 3
    for line in lines:
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = Font(name="Calibri", size=11, bold=True) if line and line[0].isupper() and not line.startswith(("-", "IMPORTANT")) and len(line) < 40 else BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1


def _build_rcm_sheet(ws: Worksheet):
    ws.freeze_panes = "A2"
    for idx, (name, width) in enumerate(zip(COLUMNS, COLUMN_WIDTHS), start=1):
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER

    for r, entry in enumerate(logic.RISK_CONTROL_MATRIX, start=2):
        fill = DOMAIN_FILL.get(entry["Domain"])
        for c, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=entry[col])
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if fill:
                cell.fill = fill
        ws.row_dimensions[r].height = 60


if __name__ == "__main__":
    build()
